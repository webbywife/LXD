"""
Module & Submodule Generator
Parses uploaded course guides (PDF/DOCX/XLSX/TXT) and generates LMS-ready content.
Each submodule is structured into 5 instructional sections:
  a. Overview       — objectives + resources
  b. Teach & Learn  — lesson content + PPTX outline
  c. Practice       — reinforcement activities
  d. Assessment     — formative quiz + authentic assessment
  e. Rubric         — topic-based scoring rubric
"""
import io
import json
import re
from typing import Optional, Tuple


def extract_text_from_file(file_content: bytes, filename: str) -> Tuple[str, str]:
    """Extract plain text from PDF, DOCX, XLSX, or TXT. Returns (text, error)."""
    ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else 'txt'
    try:
        if ext == 'pdf':
            import pdfplumber
            with pdfplumber.open(io.BytesIO(file_content)) as pdf:
                return '\n'.join(p.extract_text() or '' for p in pdf.pages), None
        elif ext == 'docx':
            import docx
            doc = docx.Document(io.BytesIO(file_content))
            return '\n'.join(p.text for p in doc.paragraphs if p.text.strip()), None
        elif ext in ('xlsx', 'xls'):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_content), read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    row_text = '\t'.join(str(c) if c is not None else '' for c in row)
                    if row_text.strip():
                        lines.append(row_text)
            return '\n'.join(lines), None
        else:
            return file_content.decode('utf-8', errors='replace'), None
    except Exception as e:
        return '', f'Could not read file: {e}'


def parse_course_guide(text: str, api_key: str) -> Tuple[Optional[dict], str]:
    """
    Use Claude to extract a module/submodule structure from course guide text.
    Returns (structure_dict, error).
    """
    if not api_key:
        return None, 'API key is required.'

    prompt = f"""You are an instructional designer. Analyze this course guide and extract a structured course outline.

COURSE GUIDE:
{text[:12000]}

Return ONLY a JSON object in this exact format (no markdown, no extra text):
{{
  "course_title": "...",
  "course_description": "1-2 sentence course summary",
  "modules": [
    {{
      "id": "M1",
      "title": "Module 1: Module Title",
      "description": "Brief module description",
      "submodules": [
        {{
          "id": "M1.1",
          "title": "Submodule 1.1: Topic Name",
          "description": "What this submodule covers",
          "topics": ["key topic 1", "key topic 2", "key topic 3"]
        }}
      ]
    }}
  ]
}}

Requirements:
- 3-8 modules total
- 2-5 submodules per module
- Derive structure directly from the course content
- If structure is implicit, infer logical groupings from the subject matter
- Use the terminology from the course guide"""

    try:
        import anthropic
        client = anthropic.Anthropic(api_key=api_key)
        msg = client.messages.create(
            model='claude-sonnet-4-6',
            max_tokens=5000,
            messages=[{'role': 'user', 'content': prompt}]
        )
        raw = msg.content[0].text.strip()
        m = re.search(r'```(?:json)?\s*([\s\S]+?)\s*```', raw)
        if m:
            raw = m.group(1)
        return json.loads(raw), None
    except json.JSONDecodeError as e:
        return None, f'AI returned invalid JSON: {e}'
    except Exception as e:
        return None, str(e)


def generate_submodule_content(
    course_title: str,
    module_title: str,
    submodule: dict,
    course_context: str,
    api_key: str,
) -> Tuple[Optional[dict], str]:
    """
    Generate all 5 instructional sections for a submodule.
    Returns (sections_dict, error) where sections_dict has keys:
      overview, teach_and_learn, practice, assessment, rubric
    """
    if not api_key:
        return None, 'API key is required.'

    topics_list = '\n'.join(f'- {t}' for t in submodule.get('topics', []))
    sub_title = submodule['title']

    prompt = f"""You are an expert instructional designer. Generate a complete 5-section educational content package for this course submodule.

Course: {course_title}
Module: {module_title}
Submodule: {sub_title}
Description: {submodule.get('description', '')}
Key Topics:
{topics_list}

Context: {course_context[:1500]}

Return ONLY a valid JSON object. All HTML values must use only double-quoted attributes. No markdown fences.

{{
  "overview": {{
    "objectives": [
      "By the end of this submodule, students will be able to [action verb] [specific outcome]"
    ],
    "resources": [
      "Textbook: Chapter X — Topic Name",
      "Reference: Title by Author / Source"
    ],
    "html": "<p>Engaging 2-sentence introduction to the submodule and its relevance.</p><h3>What You Will Learn</h3><ul><li>Learning point 1</li><li>Learning point 2</li></ul>"
  }},
  "teach_and_learn": {{
    "html": "<h2>Introduction</h2><p>Hook or opening scenario...</p><h2>Key Concepts</h2><h3>Concept 1</h3><p>Explanation with example.</p><blockquote>Key definition or principle.</blockquote><h3>Concept 2</h3><p>Explanation...</p><h2>Summary</h2><p>Wrap-up of main ideas.</p>",
    "pptx_md": "# {sub_title}\\n\\n**Course:** {course_title}\\n\\n## Learning Objectives\\n- Objective 1\\n- Objective 2\\n- Objective 3\\n\\n## Key Concepts\\n- Concept 1: brief explanation\\n- Concept 2: brief explanation\\n- Concept 3: brief explanation\\n\\n## Lesson Procedure\\n### Introduction (10 min)\\n- Engage students with a hook\\n- Connect to prior knowledge\\n\\n### Instruction (20 min)\\n- Present main concepts\\n- Show examples and non-examples\\n\\n### Guided Practice (15 min)\\n- Work through examples together\\n- Check for understanding\\n\\n## Assessment\\n- Formative check question 1\\n- Formative check question 2"
  }},
  "practice": {{
    "html": "<h2>Reinforcement Activities</h2><h3>Activity 1: [Descriptive Name]</h3><p>Clear step-by-step instructions for the activity. What students do, how long, what materials.</p><h3>Activity 2: [Descriptive Name]</h3><p>Instructions...</p><h3>Reflection Prompt</h3><p>A reflection question or journal prompt connecting the activity to the learning objectives.</p>"
  }},
  "assessment": {{
    "quiz_html": "<h2>Formative Quiz</h2><ol><li><p><strong>Question text?</strong></p><p>A. Option one<br/>B. Option two<br/>C. Option three<br/>D. Option four</p><p><em>Correct Answer: A</em></p></li></ol>",
    "quiz_questions": [
      {{
        "question": "Question text?",
        "options": ["A. Option one", "B. Option two", "C. Option three", "D. Option four"],
        "answer": "A"
      }}
    ],
    "authentic_html": "<h2>Authentic Assessment</h2><h3>The Challenge / Project</h3><p>Compelling scenario or context that makes this real and meaningful for students.</p><h3>Your Task</h3><p>Clear description of what students must do.</p><h3>Deliverables</h3><ul><li>Deliverable 1</li><li>Deliverable 2</li></ul><h3>Success Criteria</h3><p>Students will be evaluated using the rubric below. Aim for Proficient or Excellent in all criteria.</p>"
  }},
  "rubric": {{
    "title": "Assessment Rubric: {sub_title}",
    "criteria": [
      {{
        "criterion": "Criterion Name",
        "excellent": "Exceeds all expectations. Specific, observable descriptor.",
        "proficient": "Meets all expectations. Specific descriptor.",
        "developing": "Partially meets expectations. Specific descriptor.",
        "beginning": "Does not yet meet expectations. Specific descriptor."
      }}
    ],
    "html": "<table><thead><tr><th>Criterion</th><th>Excellent (4)</th><th>Proficient (3)</th><th>Developing (2)</th><th>Beginning (1)</th></tr></thead><tbody><tr><td>Criterion</td><td>Excellent descriptor</td><td>Proficient descriptor</td><td>Developing descriptor</td><td>Beginning descriptor</td></tr></tbody></table>"
  }}
}}

Specific requirements:
- objectives: exactly 4-5 SMART objectives starting with an action verb
- resources: 3-4 real-sounding references (textbooks, websites, videos) relevant to the topic
- quiz_questions: exactly 5 MCQ questions, each with 4 options (A-D), correct answer marked
- authentic assessment: choose "challenge-based" for analytical/scientific topics, "project-based" for creative/applied topics — make it compelling and realistic
- rubric: 3-5 criteria directly tied to the authentic assessment task, descriptors must be specific and observable
- All HTML: valid, no html/body/head tags, use only standard tags"""

    try:
        import anthropic
        client = anthropic.Anthropic(api_key=api_key)
        msg = client.messages.create(
            model='claude-sonnet-4-6',
            max_tokens=7000,
            messages=[{'role': 'user', 'content': prompt}]
        )
        raw = msg.content[0].text.strip()
        m = re.search(r'```(?:json)?\s*([\s\S]+?)\s*```', raw)
        if m:
            raw = m.group(1)
        data = json.loads(raw)
        # Ensure required keys exist
        for key in ('overview', 'teach_and_learn', 'practice', 'assessment', 'rubric'):
            if key not in data:
                data[key] = {}
        return data, None
    except json.JSONDecodeError as e:
        return None, f'AI returned invalid JSON: {e}'
    except Exception as e:
        return None, str(e)
